"""Utility for reading Excel (.xlsx) files."""

import openpyxl


class XlsxReader:
    """Read an xlsx workbook, returning sheet data as lists of dicts.

    Usage (context manager)::

        with XlsxReader("data.xlsx") as reader:
            rows = reader.read_sheet()          # active sheet
            rows = reader.read_sheet("Sheet2")  # by name
            all_sheets = reader.read_all_sheets()

    Usage (manual)::

        reader = XlsxReader("data.xlsx")
        reader.open()
        rows = reader.read_sheet()
        reader.close()
    """

    def __init__(self, file_path):
        self.file_path = file_path
        self._workbook = None

    # ------------------------------------------------------------------
    # Context-manager support
    # ------------------------------------------------------------------

    def __enter__(self):
        self.open()
        return self

    def __exit__(self, *args):
        self.close()

    # ------------------------------------------------------------------
    # Lifecycle
    # ------------------------------------------------------------------

    def open(self):
        """Open the workbook.  Returns *self* to allow method chaining."""
        self._workbook = openpyxl.load_workbook(self.file_path, read_only=True, data_only=True)
        return self

    def close(self):
        """Close the workbook and release resources."""
        if self._workbook is not None:
            self._workbook.close()
            self._workbook = None

    # ------------------------------------------------------------------
    # Reading helpers
    # ------------------------------------------------------------------

    def get_sheet_names(self):
        """Return a list of sheet names in the workbook."""
        self._require_open()
        return self._workbook.sheetnames

    def read_sheet(self, sheet_name=None):
        """Read *sheet_name* (or the active sheet if omitted) as a list of dicts.

        The first non-empty row is treated as the header; subsequent rows are
        returned as ``{column_header: cell_value, ...}`` mappings.  Rows that
        are entirely ``None`` are skipped.

        Args:
            sheet_name: Name of the sheet to read, or ``None`` for the active sheet.

        Returns:
            A list of dicts, one per data row.
        """
        self._require_open()
        sheet = self._workbook[sheet_name] if sheet_name else self._workbook.active
        rows = [row for row in sheet.iter_rows(values_only=True) if any(v is not None for v in row)]
        if not rows:
            return []
        headers = [str(h) if h is not None else "" for h in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:]]

    def read_all_sheets(self):
        """Read every sheet and return a ``{sheet_name: [row_dict, ...]}`` mapping."""
        return {name: self.read_sheet(name) for name in self.get_sheet_names()}

    # ------------------------------------------------------------------
    # Internal
    # ------------------------------------------------------------------

    def _require_open(self):
        if self._workbook is None:
            raise RuntimeError("Workbook is not open. Call open() or use the context manager first.")
